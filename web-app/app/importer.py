"""Lecture et génération de classeurs Excel (import des listes de prix)."""
from __future__ import annotations

import posixpath
import re
import unicodedata
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.utils import parse_price

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
REL_NS = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
HEADERS = ("N°", "Désignation", "Unité de mesure", "Prix unitaire HT")
HEADER_ALIASES = (
    {"n°", "n° article", "code", "no", "n° d article"},
    {
        "désignation",
        "désignation de l'article",
        "désignation article",
        "article",
        "libellé",
        "libelle",
    },
    {"unité", "unité de mesure", "unite", "unite de mesure"},
    {"prix unitaire ht", "prix unitaire (ht)", "prix ht", "prix unitaire"},
)


def _normalize_label(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    return " ".join(text.split()).casefold()


def _column_number(cell_ref: str) -> int:
    match = re.match(r"([A-Z]+)", cell_ref.upper())
    if not match:
        return 0
    number = 0
    for character in match.group(1):
        number = number * 26 + ord(character) - 64
    return number


def _shared_strings(workbook: ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in workbook.namelist():
        return []
    root = ET.fromstring(workbook.read("xl/sharedStrings.xml"))
    return [
        "".join(text.text or "" for text in item.iter(f"{{{NS['m']}}}t"))
        for item in root.findall("m:si", NS)
    ]


def _cell_value(cell, shared: list[str]) -> str:
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        return "".join(text.text or "" for text in cell.iter(f"{{{NS['m']}}}t"))
    value = cell.find("m:v", NS)
    if value is None:
        return ""
    text = value.text or ""
    if cell_type == "s" and text:
        try:
            return shared[int(text)]
        except (ValueError, IndexError):
            return text
    return text


def _sheet_target(workbook: ZipFile, relationship_id: str) -> str:
    relationships = ET.fromstring(workbook.read("xl/_rels/workbook.xml.rels"))
    targets = {
        relation.attrib["Id"]: relation.attrib["Target"]
        for relation in relationships.findall("r:Relationship", REL_NS)
    }
    target = targets[relationship_id].replace("\\", "/")
    if target.startswith("/"):
        return posixpath.normpath(target.lstrip("/"))
    return posixpath.normpath(posixpath.join("xl", target))


def _read_xlsx(path: str | Path) -> dict[str, list[list[str]]]:
    path = Path(path)
    with ZipFile(path) as workbook:
        shared = _shared_strings(workbook)
        root = ET.fromstring(workbook.read("xl/workbook.xml"))
        result: dict[str, list[list[str]]] = {}
        sheets = root.find("m:sheets", NS)
        if sheets is None:
            return result
        for sheet in sheets:
            name = sheet.attrib["name"]
            relationship_id = sheet.attrib[f"{{{NS['r']}}}id"]
            sheet_root = ET.fromstring(
                workbook.read(_sheet_target(workbook, relationship_id))
            )
            rows: list[list[str]] = []
            for row in sheet_root.findall(".//m:sheetData/m:row", NS):
                values: dict[int, str] = {}
                for cell in row.findall("m:c", NS):
                    values[_column_number(cell.attrib.get("r", ""))] = _cell_value(
                        cell, shared
                    )
                if values:
                    rows.append([values.get(index, "") for index in range(1, 5)])
            if not rows:
                result[name] = []
                continue
            headers = [value.strip() for value in rows[0]]
            normalized = [_normalize_label(value) for value in headers[:4]]
            if len(normalized) < 4 or not all(
                normalized[index] in aliases
                for index, aliases in enumerate(HEADER_ALIASES)
            ):
                raise ValueError(
                    f"La feuille '{name}' doit commencer par les colonnes : "
                    + ", ".join(HEADERS)
                )
            result[name] = rows[1:]
    return result


def read_sheet_names(path: str | Path) -> list[str]:
    path = Path(path)
    with ZipFile(path) as workbook:
        root = ET.fromstring(workbook.read("xl/workbook.xml"))
        sheets = root.find("m:sheets", NS)
        if sheets is None:
            return []
        return [sheet.attrib["name"] for sheet in sheets]


def import_articles(
    path: str | Path, sheet_map: dict[str, str] | None = None
) -> list[tuple[str, str, str, str, float]]:
    """Lit les lignes ``(liste, code, désignation, unité, prix)`` du classeur."""
    workbook = _read_xlsx(path)
    if sheet_map is None:
        sheet_map = {name: name for name in workbook}

    actual_names = {name.strip().lower(): name for name in workbook}
    rows: list[tuple[str, str, str, str, float]] = []
    missing: list[str] = []
    for category, sheet_name in sheet_map.items():
        actual_name = actual_names.get(str(sheet_name).strip().lower())
        if actual_name is None:
            missing.append(str(sheet_name))
            continue
        for index, row in enumerate(workbook[actual_name], start=2):
            code, designation, unit, price = [str(value).strip() for value in row[:4]]
            if not code and not designation and not unit and not price:
                continue
            if _normalize_label(code) in ("total", "totaux") or (
                _normalize_label(designation) in ("total", "totaux") and not unit
            ):
                continue
            unit = unit or "—"
            if not code or not designation:
                raise ValueError(
                    f"Ligne {index} de '{category}' : le code et la désignation "
                    "sont obligatoires."
                )
            try:
                numeric_price = parse_price(price)
            except ValueError as exc:
                raise ValueError(
                    f"Ligne {index} de '{category}' : prix invalide {price!r}."
                ) from exc
            rows.append((category, code, designation, unit, numeric_price))

    if missing:
        raise ValueError(
            "Feuille(s) introuvable(s) dans le classeur : " + ", ".join(missing)
        )
    return rows


def _sheet_titles(categories: list[str]) -> list[str]:
    titles: list[str] = []
    used: set[str] = set()
    invalid = set("[]:*?/\\")
    for category in categories or ["Liste"]:
        base = "".join(
            character for character in str(category) if character not in invalid
        ).strip()[:31] or "Liste"
        title = base
        suffix = 2
        while title.casefold() in used:
            marker = f" ({suffix})"
            title = f"{base[:31 - len(marker)]}{marker}"
            suffix += 1
        used.add(title.casefold())
        titles.append(title)
    return titles


def write_template(path: str | Path, categories: list[str]) -> Path:
    """Modèle vierge : en-têtes seuls, une feuille par liste (pas de ligne d'exemple)."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title in _sheet_titles(categories):
        sheet = workbook.create_sheet(title=title)
        sheet.append(list(HEADERS))
        for column, width in zip("ABCD", (12, 60, 18, 20)):
            sheet.column_dimensions[column].width = width
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
    workbook.save(path)
    return path
