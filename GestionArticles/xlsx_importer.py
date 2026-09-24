from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import re

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
REL_NS = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}


def _column_number(cell_ref: str) -> int:
    m = re.match(r"([A-Z]+)", cell_ref.upper())
    if not m:
        return 0
    n = 0
    for ch in m.group(1):
        n = n * 26 + ord(ch) - 64
    return n


def _shared_strings(z: ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in z.namelist():
        return []
    root = ET.fromstring(z.read("xl/sharedStrings.xml"))
    values = []
    for si in root.findall("m:si", NS):
        values.append("".join(t.text or "" for t in si.iter("{%s}t" % NS["m"])))
    return values


def _cell_value(c, shared: list[str]) -> str:
    v = c.find("m:v", NS)
    if v is None:
        return ""
    value = v.text or ""
    if c.attrib.get("t") == "s" and value:
        try:
            return shared[int(value)]
        except (ValueError, IndexError):
            return value
    return value


def read_xlsx(path: str | Path) -> dict[str, list[dict]]:
    path = Path(path)
    with ZipFile(path) as z:
        shared = _shared_strings(z)
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        relmap = {
            r.attrib["Id"]: r.attrib["Target"]
            for r in rels.findall("r:Relationship", REL_NS)
        }
        result = {}
        for sheet in wb.find("m:sheets", NS):
            name = sheet.attrib["name"]
            rid = sheet.attrib["{%s}id" % NS["r"]]
            target = relmap[rid]
            if not target.startswith("xl/"):
                target = "xl/" + target
            root = ET.fromstring(z.read(target))
            rows = []
            for row in root.findall(".//m:sheetData/m:row", NS):
                vals = {}
                for c in row.findall("m:c", NS):
                    vals[_column_number(c.attrib.get("r", ""))] = _cell_value(c, shared)
                if vals:
                    rows.append([vals.get(i, "") for i in range(1, 5)])

            if not rows:
                result[name] = []
                continue

            headers = [x.strip().lower() for x in rows[0]]
            if not all(h for h in headers[:4]):
                raise ValueError(f"La feuille '{name}' ne possède pas les 4 colonnes attendues.")

            result[name] = rows[1:]
    return result


HEADERS = ("N°", "Désignation", "Unité de mesure", "Prix unitaire HT")


def read_sheet_names(path: str | Path) -> list[str]:
    """Retourne la liste des noms de feuilles sans valider leur contenu."""
    path = Path(path)
    with ZipFile(path) as z:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        return [sheet.attrib["name"] for sheet in wb.find("m:sheets", NS)]


def parse_price(price: str) -> float:
    """Normalise les formats de prix (virgule, séparateurs de milliers)."""
    numeric_text = str(price).replace("\u00a0", "").replace(" ", "")
    # Accepte 1234.56, 1 234,56, 1,234.56 et même 92,028,48.
    if "," in numeric_text and "." in numeric_text:
        if numeric_text.rfind(",") > numeric_text.rfind("."):
            numeric_text = numeric_text.replace(".", "").replace(",", ".")
        else:
            numeric_text = numeric_text.replace(",", "")
    elif numeric_text.count(",") > 1:
        parts = numeric_text.split(",")
        numeric_text = "".join(parts[:-1]) + "." + parts[-1]
    elif numeric_text.count(",") == 1:
        left, right = numeric_text.split(",")
        numeric_text = left + "." + right
    return float(numeric_text)


def extract_articles(
    path: str | Path,
    sheet_map: dict[str, str] | None = None,
) -> list[tuple[str, str, str, str, float]]:
    """Extrait les articles d'un classeur.

    `sheet_map` associe un nom de liste (catégorie) au nom de la feuille Excel
    correspondante. Sans mapping, chaque feuille devient sa propre liste.
    """
    workbook = read_xlsx(path)

    if sheet_map is None:
        sheet_map = {name: name for name in workbook}

    lower_names = {name.strip().lower(): name for name in workbook}
    out: list[tuple[str, str, str, str, float]] = []
    missing: list[str] = []

    for category, sheet_name in sheet_map.items():
        actual = lower_names.get(str(sheet_name).strip().lower())
        if actual is None:
            missing.append(str(sheet_name))
            continue
        for row in workbook[actual]:
            if len(row) < 4:
                continue
            code, designation, unit, price = [str(v).strip() for v in row[:4]]
            if not code and not designation:
                continue
            # Certaines feuilles comportent une ligne de total en fin de liste.
            if code.upper() == "TOTAL" and not price:
                continue
            try:
                numeric_price = parse_price(price)
            except ValueError:
                raise ValueError(f"Prix invalide dans '{category}': {price!r}")
            out.append((category, code, designation, unit, numeric_price))

    if missing:
        raise ValueError("Feuille(s) introuvable(s) dans le classeur : " + ", ".join(missing))
    return out


def write_template(path: str | Path, categories: list[str]) -> Path:
    """Crée un classeur modèle vierge, une feuille par liste."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    invalid = set(r'[]:*?/\'')
    for category in categories or ["Liste"]:
        title = "".join(ch for ch in str(category) if ch not in invalid).strip()[:31] or "Liste"
        ws = wb.create_sheet(title=title)
        ws.append(list(HEADERS))
        for column, width in zip("ABCD", (12, 60, 18, 20)):
            ws.column_dimensions[column].width = width
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
        ws.append([1, "Exemple d'article (à supprimer)", "pièce", 1000])
    wb.save(path)
    return path
