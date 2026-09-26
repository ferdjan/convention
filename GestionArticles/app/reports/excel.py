from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.utils import excel_safe_text

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
LABEL_FONT = Font(bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="EAF2F8")
MONEY_FORMAT = '#,##0.00" DA"'
THIN = Side(style="thin", color="B7C9D6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = ("N°", "Désignation", "Unité", "Prix unitaire HT", "Quantité", "Total HT")
WIDTHS = (10, 55, 16, 20, 12, 18)


def build_excel(
    path: str | Path,
    number: str,
    date_text: str,
    category: str,
    items: list[dict],
    total_ht: float,
    tva_rate: float,
    tva_amount: float | None = None,
    total_ttc: float | None = None,
) -> Path:
    if tva_amount is None:
        tva_amount = round(float(total_ht) * tva_rate, 2)
    if total_ttc is None:
        total_ttc = round(float(total_ht) + tva_amount, 2)
    tva_label = f"TVA {tva_rate * 100:g}%"

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Document"

    ws.merge_cells("A1:F1")
    ws["A1"] = "DOCUMENT DE COMMANDE"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A3"] = "N° :"
    ws["A3"].font = LABEL_FONT
    ws["B3"] = excel_safe_text(number)
    ws["E3"] = "Date :"
    ws["E3"].font = LABEL_FONT
    ws["F3"] = excel_safe_text(date_text)
    ws["A4"] = "Liste :"
    ws["A4"].font = LABEL_FONT
    ws["B4"] = excel_safe_text(category)

    header_row = 6
    for col, title in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = header_row + 1
    for item in items:
        quantity = float(item["quantity"])
        unit_price = float(item["unit_price_ht"])
        ws.cell(row=row_idx, column=1, value=excel_safe_text(item["code"]))
        ws.cell(row=row_idx, column=2, value=excel_safe_text(item["designation"]))
        ws.cell(row=row_idx, column=3, value=excel_safe_text(item["unit"]))
        ws.cell(row=row_idx, column=4, value=unit_price).number_format = MONEY_FORMAT
        ws.cell(row=row_idx, column=5, value=quantity)
        ws.cell(row=row_idx, column=6, value=round(quantity * unit_price, 2)).number_format = MONEY_FORMAT
        row_idx += 1

    ht_row = row_idx
    tva_row = row_idx + 1
    ttc_row = row_idx + 2
    for row, label, value in (
        (ht_row, "TOTAL HT", round(float(total_ht), 2)),
        (tva_row, tva_label, tva_amount),
        (ttc_row, "TOTAL TTC", total_ttc),
    ):
        ws.cell(row=row, column=5, value=label).font = LABEL_FONT
        cell = ws.cell(row=row, column=6, value=value)
        cell.font = LABEL_FONT
        cell.number_format = MONEY_FORMAT

    for row in ws.iter_rows(min_row=header_row, max_row=ttc_row, min_col=1, max_col=6):
        for cell in row:
            cell.border = BORDER
            if cell.column in (1, 3, 5):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif cell.column == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
    for col in (5, 6):
        ws.cell(row=ht_row, column=col).fill = TOTAL_FILL
        ws.cell(row=tva_row, column=col).fill = TOTAL_FILL
        ws.cell(row=ttc_row, column=col).fill = HEADER_FILL
        ws.cell(row=ttc_row, column=col).font = HEADER_FONT

    for col, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.cell(row=ttc_row + 2, column=1, value="Document généré par l'application de gestion des articles.")
    wb.save(path)
    return path
